"""石油製品価格調査 給油所小売価格（週次）の取得・整形。

資源エネルギー庁が公表する石油製品価格調査の「週次ファイル」（1990年8月27日以降の
全調査日を収めた Excel）を取得し、調査日×油種×地域の 1 行 = 1 レコードへ整形して
CSV に保存する。

ファイル名は公表のたびに変わる（例 261007s5.xlsx）ので、URL は組み立てずに
調査結果のページから拾う。

シートは油種ごと（ハイオク・レギュラー・軽油・灯油店頭・灯油配達）で、調査日が縦、
地域（全国・経済産業局・都道府県）が横に並ぶ。都道府県の列は 2004年6月14日以降にしか
値が無い。ガソリンと軽油のシートには価格に含まれる税（ガソリン税・軽油引取税）の列が、
全シートに消費税率の列がある。
"""

import csv
import logging
import re
from datetime import datetime
from pathlib import Path

import openpyxl

from enecho import fetch_file, fetch_text
from power_survey_common import PREFECTURE_CODES, squash

logger = logging.getLogger("pipelines")

RESULTS_PATH = "/statistics/petroleum_and_lpgas/pl007/results.html"
WEEKLY_FILE_PATTERN = re.compile(
    r"/statistics/petroleum_and_lpgas/pl007/xlsx/\d{6}s5\.xlsx"
)

# シート名と単位。単位は見出しの「（\/ﾘｯﾄﾙ）」「（\/18ﾘｯﾄﾙ）」でも確かめる。
PRODUCTS = {
    "ハイオク": "円/L",
    "レギュラー": "円/L",
    "軽油": "円/L",
    "灯油店頭": "円/18L",
    "灯油配達": "円/18L",
}

DATE_HEADING = "調査日"
# 価格に含まれる税の見出し。灯油のシートには無い。
TAX_HEADINGS = {
    "ハイオク": "（うちガソリン税）",
    "レギュラー": "（うちガソリン税）",
    "軽油": "(うち軽油引取税）",
}
CONSUMPTION_TAX_HEADING = "消費税率"

# 地域の見出しは都道府県名を「県」「府」「都」抜きで書き、経済産業局は「局」で終わる。
# 北海道局と沖縄局はそれぞれ1道県だけを管轄するので、都道府県コードも振る。
BUREAU_PREFECTURES = {"北海道局": "北海道", "沖縄局": "沖縄県"}
_PREFECTURE_BY_SHORT_NAME = {
    re.sub(r"[都府県]$", "", name): name
    for name in PREFECTURE_CODES
    if name not in BUREAU_PREFECTURES.values()
}

OUTPUT_COLUMNS = [
    "survey_date",
    "product_name",
    "price_unit",
    "area_order",
    "area_name",
    "area_level",
    "pref_code",
    "pref_name",
    "price_yen",
    "fuel_tax_yen",
    "consumption_tax_rate",
]


def resolve_source() -> str:
    """調査結果のページから週次ファイルの URL を拾う。"""
    found = set(WEEKLY_FILE_PATTERN.findall(fetch_text(RESULTS_PATH)))
    if len(found) != 1:
        raise RuntimeError(f"週次ファイルを1つに絞れない: {sorted(found)}")
    return found.pop()


def classify_area(label: str) -> tuple[str, str | None, str | None]:
    """地域の見出しから（区分, 都道府県コード, 都道府県名）を決める。"""
    if label == "全国":
        return "全国", None, None
    if label in BUREAU_PREFECTURES:
        pref_name = BUREAU_PREFECTURES[label]
        return "経済産業局", PREFECTURE_CODES[pref_name], pref_name
    if label.endswith("局"):
        return "経済産業局", None, None
    pref_name = _PREFECTURE_BY_SHORT_NAME.get(label)
    if pref_name is None:
        raise RuntimeError(f"想定外の地域の見出し: {label}")
    return "都道府県", PREFECTURE_CODES[pref_name], pref_name


def parse_sheet(ws, product_name: str) -> list[dict]:
    rows = ws.iter_rows(values_only=True)
    header = [squash(cell) for cell in next(rows)]

    unit = PRODUCTS[product_name]
    expected_unit = "/18ﾘｯﾄﾙ" if unit == "円/18L" else "/ﾘｯﾄﾙ"
    if not header[0].startswith(product_name) or expected_unit not in header[0]:
        raise RuntimeError(f"{product_name}: 想定外の見出し {header[0]!r}")
    if header[1] != DATE_HEADING:
        raise RuntimeError(f"{product_name}: 2列目が調査日ではない {header[1]!r}")

    # 地域の列は空の見出しの手前まで。その後ろに税の列が続く。
    end = header.index("", 2)
    areas = [
        (index, header[index], *classify_area(header[index])) for index in range(2, end)
    ]
    if header[2] != "全国":
        raise RuntimeError(f"{product_name}: 3列目が全国ではない {header[2]!r}")

    tax_index = None
    consumption_tax_index = None
    for index in range(end + 1, len(header)):
        if header[index] == TAX_HEADINGS.get(product_name):
            tax_index = index
        elif header[index] == CONSUMPTION_TAX_HEADING:
            consumption_tax_index = index
        elif header[index]:
            raise RuntimeError(f"{product_name}: 想定外の列 {header[index]!r}")
    if product_name in TAX_HEADINGS and tax_index is None:
        raise RuntimeError(f"{product_name}: 税の列が無い")
    if consumption_tax_index is None:
        raise RuntimeError(f"{product_name}: 消費税率の列が無い")

    records = []
    for row in rows:
        survey_date = row[1]
        if survey_date is None:
            # 末尾の注記行（消費税の扱い）。値は持たない。
            if any(isinstance(value, (int, float)) for value in row[2:]):
                raise RuntimeError(
                    f"{product_name}: 調査日の無い行に値がある {row[:4]}"
                )
            continue
        if not isinstance(survey_date, datetime):
            raise TypeError(f"{product_name}: 調査日が日付ではない {survey_date!r}")
        fuel_tax = row[tax_index] if tax_index is not None else None
        consumption_tax = row[consumption_tax_index]
        for order, (index, label, level, pref_code, pref_name) in enumerate(areas, 1):
            value = row[index]
            if value is None:
                continue
            if not isinstance(value, (int, float)):
                raise TypeError(f"{product_name}: 数値ではない価格 {label} {value!r}")
            records.append(
                {
                    "survey_date": survey_date.date().isoformat(),
                    "product_name": product_name,
                    "price_unit": unit,
                    "area_order": order,
                    "area_name": label,
                    "area_level": level,
                    "pref_code": pref_code,
                    "pref_name": pref_name,
                    "price_yen": value,
                    "fuel_tax_yen": fuel_tax,
                    "consumption_tax_rate": consumption_tax,
                }
            )
    return records


def parse(xlsx_path: Path) -> list[dict]:
    wb = openpyxl.load_workbook(xlsx_path, read_only=True, data_only=True)
    if wb.sheetnames != list(PRODUCTS):
        raise RuntimeError(f"想定外のシート構成: {wb.sheetnames}")
    records = []
    for product_name in PRODUCTS:
        records.extend(parse_sheet(wb[product_name], product_name))
    return records


def download_and_parse(csv_path: Path) -> int:
    source = resolve_source()
    xlsx_path = csv_path.parent / "petroleum" / Path(source).name
    fetch_file(source, xlsx_path)
    records = parse(xlsx_path)
    with csv_path.open("w", newline="", encoding="utf-8") as f:
        writer = csv.DictWriter(f, fieldnames=OUTPUT_COLUMNS)
        writer.writeheader()
        writer.writerows(records)
    return len(records)
